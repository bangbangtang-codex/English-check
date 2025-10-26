"""Import Excel/CSV files into the database."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Optional

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional dependency
    load_workbook = None  # type: ignore[assignment]

from . import database
from .utils import generate_id, meaning_key, normalize_term, parse_tags, timestamp_now


@dataclass
class ImportRow:
    term: str
    cn: Optional[str]
    ipa: Optional[str]
    tags: List[str]
    notes: Optional[str]
    source_file: str


@dataclass
class ImportResult:
    created: int
    updated: int
    skipped: int
    conflicts: int
    conflict_rows: List[ImportRow]


def parse_file(path: str) -> List[ImportRow]:
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(path)
    suffix = file_path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xltx"}:
        if load_workbook is None:
            raise RuntimeError(
                "openpyxl is required to import Excel files. Install dependencies with "
                "`pip install -r requirements.txt`."
            )
        return _parse_excel(file_path)
    if suffix in {".csv", ".tsv"}:
        return _parse_csv(file_path)
    raise ValueError(f"Unsupported file type: {file_path.suffix}")


def _parse_excel(path: Path) -> List[ImportRow]:
    assert load_workbook is not None
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb.active
    rows: List[ImportRow] = []
    header_processed = False
    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        term = (row[0] or "").strip()
        cn = (row[1] or "").strip() or None
        ipa = (row[2] or "").strip() or None
        notes = (row[3] or "").strip() or None
        tags = parse_tags(row[4] if len(row) > 4 else None)
        if not header_processed and _is_header(term, cn, ipa):
            header_processed = True
            continue
        if not term:
            continue
        rows.append(
            ImportRow(
                term=term,
                cn=cn,
                ipa=ipa,
                tags=tags,
                notes=notes,
                source_file=path.name,
            )
        )
    return rows


def _parse_csv(path: Path) -> List[ImportRow]:
    rows: List[ImportRow] = []
    delimiter = "," if path.suffix.lower() == ".csv" else "\t"
    with path.open("r", encoding="utf-8") as fh:
        reader = csv.reader(fh, delimiter=delimiter)
        header_processed = False
        for line in reader:
            if not line:
                continue
            term = (line[0] if len(line) > 0 else "").strip()
            cn = (line[1] if len(line) > 1 else "").strip() or None
            ipa = (line[2] if len(line) > 2 else "").strip() or None
            notes = (line[3] if len(line) > 3 else "").strip() or None
            tags = parse_tags(line[4] if len(line) > 4 else None)
            if not header_processed and _is_header(term, cn, ipa):
                header_processed = True
                continue
            if not term:
                continue
            rows.append(
                ImportRow(
                    term=term,
                    cn=cn,
                    ipa=ipa,
                    tags=tags,
                    notes=notes,
                    source_file=path.name,
                )
            )
    return rows


def _is_header(term: str, cn: Optional[str], ipa: Optional[str]) -> bool:
    header_tokens = {"word", "english", "term", "中文", "释义", "meaning", "ipa", "音标"}
    combined = " ".join(filter(None, [term.lower(), (cn or "").lower(), (ipa or "").lower()]))
    return any(token in combined for token in header_tokens)


def import_rows(rows: Iterable[ImportRow], *, source_name: str = "manual") -> ImportResult:
    stats = {"created": 0, "updated": 0, "skipped": 0, "conflicts": 0}
    conflict_rows: List[ImportRow] = []
    now = timestamp_now()
    with database.connect() as conn:
        for row in rows:
            normalized = normalize_term(row.term)
            mkey = meaning_key(row.term, row.cn)
            existing = conn.execute(
                "SELECT * FROM cards WHERE normalized_term=? AND meaning_key=?",
                (normalized, mkey),
            ).fetchone()

            card_id = generate_id(row.term, row.cn, row.ipa)
            tags = row.tags
            notes = row.notes
            if existing:
                stats["updated"] += 1
                created_at = datetime.strptime(existing["created_at"], database.ISO_FMT)
                card_id = existing["id"]
            else:
                stats["created"] += 1
                created_at = now

            card = database.Card(
                id=card_id,
                term=row.term.strip(),
                cn=row.cn,
                ipa=row.ipa,
                tags=tags,
                source_file=row.source_file,
                created_at=created_at,
                updated_at=now,
                notes=notes,
                normalized_term=normalized,
                meaning_key=mkey,
            )
            database.upsert_card(conn, card)
            database.ensure_review(conn, card_id, next_review=now)

        summary = f"created {stats['created']}, updated {stats['updated']}"
        database.record_import(conn, source_name, summary, stats)
    return ImportResult(
        created=stats["created"],
        updated=stats["updated"],
        skipped=stats["skipped"],
        conflicts=stats["conflicts"],
        conflict_rows=conflict_rows,
    )


def import_file(path: str) -> ImportResult:
    rows = parse_file(path)
    return import_rows(rows, source_name=Path(path).name)
